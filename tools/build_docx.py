#!/usr/bin/env python3
"""Build printable Word files for The Fun Drawer.

Writes print/letter/*.docx and print/a4/*.docx:
  <category>--<activity>.docx   one activity
  <category>--all.docx          a whole category booklet
  fun-drawer-everything.docx    everything

Needs: pandoc, python-docx.   Run:  python3 tools/build_docx.py
Category pictures come from img/png/*.png (make them with: node tools/render_png.js).
"""
import json
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
CATALOG = json.loads((ROOT / "catalog.json").read_text())

PAPERS = {
    "letter": dict(w=Inches(8.5), h=Inches(11), margin=Inches(0.6)),
    "a4": dict(w=Mm(210), h=Mm(297), margin=Mm(15)),
}
INK = RGBColor(0x1F, 0x2A, 0x44)
BOX_HEIGHT = {"small": 1.6, "": 2.6, "big": 4.6}  # inches

DRAW_BOX = re.compile(r'<div class="draw-box(?: (\w+))?"(?: data-label="([^"]*)")?></div>')
PANELS = re.compile(r'<div class="panels">\s*((?:<div class="draw-box[^>]*></div>\s*)+)</div>')


# ---------- markdown prep ----------

def facts_line(act):
    time = act.get("time") or f'{act["minutes"]} min'
    cash = "$" * act["cleanup"]
    return (f'⏱️ **{time}** · 🧽 Cleanup **{cash}** ({CATALOG["cleanup"][str(act["cleanup"])]})'
            f' · 👧 Ages **{act["ages"]}** · {act["where"]}')


def prep(text):
    """Swap web-only HTML for placeholders that post-processing turns into boxes."""
    def panels(m):
        labels = [lbl or "" for _, lbl in DRAW_BOX.findall(m.group(1))]
        return "DRAWGRID::" + "|".join(labels)
    text = PANELS.sub(panels, text)
    return DRAW_BOX.sub(lambda m: f"DRAWBOX::{m.group(1) or ''}::{m.group(2) or ''}", text)


def activity_md(cat, act):
    text = (ROOT / "activities" / cat["id"] / f'{act["id"]}.md').read_text()
    first, _, rest = text.partition("\n")
    return f"{first}\n\n{facts_line(act)}\n{rest}"


def category_md(cat):
    rows = "\n".join(
        f'| {a["emoji"]} {a["name"]} | ⏱️ {a.get("time") or str(a["minutes"]) + " min"} '
        f'| {"$" * a["cleanup"]} | {a["ages"]} |' for a in cat["activities"])
    cover = (f'# {cat["emoji"]} {cat["name"]}\n\n'
             f'IMAGE::img/png/{cat["id"]}.png\n\n'
             f'**{cat["tagline"]}**\n\n{cat["intro"]}\n\n'
             f'## 📂 What\'s inside\n\n| Activity | Time | 🧽 Cleanup | Ages |\n|---|---|---|---|\n{rows}\n')
    return "\n\n".join([cover] + [activity_md(cat, a) for a in cat["activities"]])


def everything_md():
    legend = "\n".join(f'- **{"$" * int(k)}** = {v}' for k, v in CATALOG["cleanup"].items())
    cover = (f'# 🗄️ {CATALOG["title"]}\n\nPick a thing from the drawer. Find a fun thing to do! 🎉\n\n'
             + "\n".join(f'- {c["emoji"]} **{c["name"]}** — {c["tagline"]}' for c in CATALOG["categories"])
             + f'\n\n### 🧽 Cleanup scale\n{legend}\n')
    return "\n\n".join([cover] + [category_md(c) for c in CATALOG["categories"]])


# ---------- docx helpers ----------

def set_borders(tbl, color="9AA3B2", size=8, style="single"):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), style)
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        borders.append(el)
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    tblPr.append(borders)


def shade(cell, fill="F3EEFF"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def box_table(doc, para, labels, cols, height):
    rows = (len(labels) + cols - 1) // cols
    tbl = doc.add_table(rows=rows, cols=cols)
    set_borders(tbl, color="B8B0A2", size=12, style="dashed")
    for i, label in enumerate(labels):
        cell = tbl.cell(i // cols, i % cols)
        run = cell.paragraphs[0].add_run(label)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x55, 0x60, 0x7A)
        run.bold = True
    for row in tbl.rows:
        row.height = Inches(height)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    para._p.addprevious(tbl._tbl)
    para._p.getparent().remove(para._p)


def reference_doc(paper, path):
    subprocess.run(["pandoc", "-o", str(path), "--print-default-data-file", "reference.docx"], check=True)
    doc = Document(str(path))
    p = PAPERS[paper]
    for s in doc.sections:
        s.page_width, s.page_height = p["w"], p["h"]
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = p["margin"]
    sizes = {"Normal": 12.5, "Body Text": 12.5, "First Paragraph": 12.5, "Compact": 12.5,
             "Title": 28, "Heading 1": 26, "Heading 2": 18, "Heading 3": 14}
    colors = {"Heading 1": INK, "Heading 2": RGBColor(0xE8, 0x61, 0x2B), "Heading 3": RGBColor(0x4B, 0x3F, 0xB8)}
    for name, size in sizes.items():
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.color.rgb = colors.get(name, INK)
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), "Arial")
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(a), None)
        if name.startswith("Heading"):
            st.font.bold = True
            st.paragraph_format.space_before = Pt(14 if name != "Heading 3" else 10)
            st.paragraph_format.space_after = Pt(4)
    doc.save(str(path))


def finish(path):
    doc = Document(str(path))
    for tbl in doc.tables:
        set_borders(tbl)
        for cell in tbl.rows[0].cells:
            shade(cell)
    first_h1 = True
    for para in list(doc.paragraphs):
        text = para.text.strip()
        style = para.style.name if para.style is not None else ""
        if style == "Heading 1":
            if not first_h1:
                para.paragraph_format.page_break_before = True
            first_h1 = False
        elif style == "Heading 2" and "Printable" in text:
            para.paragraph_format.page_break_before = True
        elif text.startswith("DRAWBOX::"):
            _, size, label = text.split("::", 2)
            box_table(doc, para, [label], 1, BOX_HEIGHT.get(size, 2.6))
        elif text.startswith("DRAWGRID::"):
            box_table(doc, para, text.split("::", 1)[1].split("|"), 2, 3.9)
        elif text.startswith("IMAGE::"):
            img = ROOT / text.split("::", 1)[1]
            for r in list(para.runs):
                r._r.getparent().remove(r._r)
            if img.exists():
                para.add_run().add_picture(str(img), width=Inches(2.4))
    doc.save(str(path))


def build(md_text, out, ref):
    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False) as f:
        f.write(prep(md_text))
    subprocess.run(["pandoc", f.name, "-f", "gfm+hard_line_breaks", "-t", "docx",
                    "--reference-doc", str(ref), "-o", str(out)], check=True)
    Path(f.name).unlink()
    finish(out)


def main():
    with tempfile.TemporaryDirectory() as tmp:
        for paper in PAPERS:
            ref = Path(tmp) / f"ref-{paper}.docx"
            reference_doc(paper, ref)
            outdir = ROOT / "print" / paper
            for cat in CATALOG["categories"]:
                for act in cat["activities"]:
                    build(activity_md(cat, act), outdir / f'{cat["id"]}--{act["id"]}.docx', ref)
                build(category_md(cat), outdir / f'{cat["id"]}--all.docx', ref)
            build(everything_md(), outdir / "fun-drawer-everything.docx", ref)
            print(f"✅ {paper}: {len(list(outdir.glob('*.docx')))} files")


if __name__ == "__main__":
    main()
